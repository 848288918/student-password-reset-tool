# -*- coding: utf-8 -*-
"""
学员密码批量重置工具
读取同目录 info.txt 中的证件号（每行一个），登录医博士教学管理中心后台，
逐个搜索学员并重置密码为 Aa123456@，结果输出到同目录 Excel 表格。
"""
import os
import sys
import json
import base64
import io
import time
import datetime

import requests
from PIL import Image
from Crypto.Cipher import AES
from Crypto.Util.Padding import pad
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ============ 配置 ============
LOGIN_URL = "https://apicloud.yiboshi.com/openplatfrom-authserver/v2/authentication/admin/login"
CAPTCHA_GET = "https://apicloud.yiboshi.com/openplatfrom-authserver/captcha/get"
CAPTCHA_CHECK = "https://apicloud.yiboshi.com/openplatfrom-authserver/captcha/check"
SEARCH_URL = "https://apiadmin.yiboshi.com/api/study/student/getStudentList"
RESET_URL = "https://apicloud.yiboshi.com/openplatfrom-userserver/user/editPasswordByUuid"

ADMIN_USER = "linhs"
# 新密码 Aa123456@ 的散列值（与后台前端一致）
NEW_PASSWORD_MD5 = "4f42142df925126c6bdfd337b6716cba"
NEW_PASSWORD_HASH2 = "2bf6f500ff4d220822762953358964016025c45bcd2e86a87c5d36402f321e46"

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")

BASE_DIR = (os.path.dirname(sys.executable) if getattr(sys, "frozen", False)
            else os.path.dirname(os.path.abspath(__file__)))
INPUT_FILE = os.path.join(BASE_DIR, "info.txt")


def aes_enc(text, key):
    c = AES.new(key.encode(), AES.MODE_ECB)
    return base64.b64encode(c.encrypt(pad(text.encode(), 16))).decode()


def find_gap_x(bg_b64, piece_b64):
    """纯 Python 滑块缺口识别（蒙版归一化互相关），返回原图 x 坐标与置信度"""
    bg = Image.open(io.BytesIO(base64.b64decode(bg_b64))).convert("L")
    piece = Image.open(io.BytesIO(base64.b64decode(piece_b64))).convert("RGBA")
    W, H = bg.size
    w, h = piece.size
    bgp = bg.load()
    pcp = piece.load()
    pts = [(x, y) for y in range(h) for x in range(w) if pcp[x, y][3] > 10]
    tv = [sum(pcp[x, y][:3]) / 3.0 for x, y in pts]
    tmean = sum(tv) / len(tv)
    tvar = sum((v - tmean) ** 2 for v in tv)
    best_x, best_score = 0, -2.0
    for x0 in range(0, W - w + 1):
        vals = [bgp[x0 + x, y] for x, y in pts]
        n = len(vals)
        mean = sum(vals) / n
        num = 0.0
        var = 0.0
        for v, t in zip(vals, tv):
            dv = v - mean
            num += dv * (t - tmean)
            var += dv * dv
        denom = (var * tvar) ** 0.5
        score = num / denom if denom else -1.0
        if score > best_score:
            best_score, best_x = score, x0
    return best_x, best_score


def login(session):
    """登录（含滑块验证码自动识别），成功返回 token，失败返回 None"""
    for attempt in range(1, 4):
        try:
            g = session.post(CAPTCHA_GET, json={"captchaType": "blockPuzzle"}, timeout=20).json()["repData"]
            gx, conf = find_gap_x(g["originalImageBase64"], g["jigsawImageBase64"])
            pj = json.dumps({"x": float(gx), "y": 5}, separators=(",", ":"))
            chk = session.post(CAPTCHA_CHECK, json={
                "captchaType": "blockPuzzle",
                "pointJson": aes_enc(pj, g["secretKey"]),
                "token": g["token"],
            }, timeout=20).json()
            if not chk["repData"]["result"]:
                print(f"  滑块验证未通过(第{attempt}次)，重试...")
                continue
            cvf = aes_enc(g["token"] + "---" + pj, g["secretKey"])
            lj = session.post(LOGIN_URL, data={
                "username": ADMIN_USER,
                "password": NEW_PASSWORD_MD5,
                "password2": NEW_PASSWORD_HASH2,
                "captchaVerification": cvf,
            }, timeout=20).json()
            if lj.get("code") == 0:
                return lj["data"]
            print(f"  登录失败(第{attempt}次): {lj.get('msg')}")
        except Exception as e:
            print(f"  登录异常(第{attempt}次): {e}")
        time.sleep(2)
    return None


def search_student(session, idcard):
    r = session.get(SEARCH_URL, params={
        "currentPage": 1, "pageSize": 10, "total": 1, "mobileHidden": 0, "sertIdHidden": 0,
        "startNum": "", "endNum": "", "id": "", "startDate": "", "endDate": "",
        "lable": "证件号", "value": idcard, "userName": "", "personType": "",
        "specialtyId1": "", "specialtyId2": "", "inputType": "", "companyName": "",
        "province": "", "city": "", "county": "", "rural": "", "isAccess": "true",
        "isDel": "", "countryDoctor": "",
    }, timeout=20)
    return r.json()


def reset_password(session, uuid):
    r = session.post(RESET_URL, json={
        "password": NEW_PASSWORD_MD5,
        "password2": NEW_PASSWORD_HASH2,
        "uuid": uuid,
    }, timeout=20)
    return r.json()


def main():
    print("=" * 50)
    print("  学员密码批量重置工具")
    print("=" * 50)
    if not os.path.exists(INPUT_FILE):
        print(f"\n未找到名单文件: {INPUT_FILE}")
        print("请在 exe 同目录创建 info.txt，每行一个证件号。")
        try:
            input("\n按回车键退出...")
        except EOFError:
            pass
        return

    with open(INPUT_FILE, "r", encoding="utf-8") as f:
        ids = [ln.strip() for ln in f if ln.strip()]
    print(f"共读取 {len(ids)} 个证件号\n")

    session = requests.Session()
    session.headers.update({"User-Agent": UA,
                            "Origin": "https://teachadmin.yiboshi.com",
                            "Referer": "https://teachadmin.yiboshi.com/"})
    print("正在登录管理后台...")
    token = login(session)
    if not token:
        print("登录失败，程序结束。")
        try:
            input("\n按回车键退出...")
        except EOFError:
            pass
        return
    session.headers.update({"Authorization": "Bearer " + token})
    print("登录成功，开始处理...\n")

    results = []
    for idx, idcard in enumerate(ids, 1):
        name = username = ""
        note = ""
        try:
            q = search_student(session, idcard)
            if q.get("status") == 401 or q.get("code") == 401:
                token = login(session)
                session.headers.update({"Authorization": "Bearer " + token})
                q = search_student(session, idcard)
            lst = (q.get("data") or {}).get("list") or []
            if not lst:
                status = "未找到学员"
            else:
                stu = lst[0]
                name = stu.get("realName", "")
                username = stu.get("userName", "") or stu.get("mobile", "")
                rp = reset_password(session, stu["uuid"])
                if rp.get("code") == 0:
                    status = "重置成功"
                else:
                    status = "重置失败"
                    note = str(rp.get("msg", ""))
        except Exception as e:
            status = "处理异常"
            note = str(e)[:100]
        results.append([idcard, name, username, status, note])
        print(f"[{idx}/{len(ids)}] {idcard}  {name}  -> {status} {note}")
        time.sleep(0.5)

    # 输出 Excel
    out_name = "重置结果_" + datetime.datetime.now().strftime("%Y%m%d_%H%M%S") + ".xlsx"
    out_path = os.path.join(BASE_DIR, out_name)
    wb = Workbook()
    ws = wb.active
    ws.title = "重置结果"
    headers = ["证件号", "姓名", "用户名", "结果", "备注"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(horizontal="center")
    ok_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    for row in results:
        ws.append(row)
        cell = ws.cell(row=ws.max_row, column=4)
        if row[3] == "重置成功":
            cell.fill = ok_fill
        else:
            cell.fill = fail_fill
    for col, width in zip("ABCDE", (22, 10, 20, 12, 30)):
        ws.column_dimensions[col].width = width
    wb.save(out_path)

    ok_cnt = sum(1 for r in results if r[3] == "重置成功")
    print("\n" + "=" * 50)
    print(f"处理完成: 成功 {ok_cnt} / 共 {len(results)}")
    print(f"结果已保存: {out_path}")
    try:
        input("\n按回车键退出...")
    except EOFError:
        pass


if __name__ == "__main__":
    main()
